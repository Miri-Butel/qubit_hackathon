"""Render the generated .pptx back out as HTML for visual QA.

Reads the actual file python-pptx produced -- not the code that made it -- so
what gets inspected is what will be imported into Slides. 96 px per inch, and
pt sizes scale by 96/72, so the browser's Arial metrics stand in for
PowerPoint's closely enough to catch overflow and collisions.
"""
import base64, sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Emu

SRC = Path("/Users/echaiezra/GitHub/school/Classiq/hackathon/qubit_hackathon/presentation/dist/qubit_pitch.pptx")
OUT = Path(sys.argv[1] if len(sys.argv) > 1 else "/tmp/preview.html")
PX = 96.0

def inches(v): return Emu(v).inches if v is not None else 0.0

prs = Presentation(str(SRC))
sw, sh = inches(prs.slide_width) * PX, inches(prs.slide_height) * PX
parts = ["""<meta charset="utf-8"><style>
body{background:#1a1a1a;margin:0;font-family:Arial,Helvetica,sans-serif}
.slide{position:relative;overflow:hidden;margin:14px auto;box-shadow:0 2px 18px #000}
.n{position:absolute;left:-11px;top:-2px;color:#888;font-size:11px}
.sh{position:absolute;box-sizing:border-box}
.tb{position:absolute;box-sizing:border-box;white-space:pre-wrap}
.ov{outline:2px solid #ff00ff !important}
</style>"""]

problems = []
for idx, slide in enumerate(prs.slides, 1):
    bg = "#060B18"
    parts.append(f'<div class="n">{idx}</div>')
    parts.append(f'<div class="slide" style="width:{sw:.0f}px;height:{sh:.0f}px;background:{bg}">')
    for shape in slide.shapes:
        x, y = inches(shape.left) * PX, inches(shape.top) * PX
        w, h = inches(shape.width) * PX, inches(shape.height) * PX
        oob = x < -1 or y < -1 or x + w > sw + 1 or y + h > sh + 1
        if oob:
            problems.append(f"slide {idx}: shape out of bounds "
                            f"({x:.0f},{y:.0f},{w:.0f}x{h:.0f})")
        cls = "ov" if oob else ""
        if shape.shape_type == 13 or shape.__class__.__name__ == "Picture":
            b64 = base64.b64encode(shape.image.blob).decode()
            parts.append(f'<img class="sh {cls}" src="data:image/png;base64,{b64}" '
                         f'style="left:{x:.1f}px;top:{y:.1f}px;width:{w:.1f}px;height:{h:.1f}px">')
            continue
        style = f"left:{x:.1f}px;top:{y:.1f}px;width:{w:.1f}px;height:{h:.1f}px;"
        if shape.has_text_frame and shape.text_frame.text.strip():
            first = None
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    first = run; break
                if first: break
            size = first.font.size.pt if first and first.font.size else 15
            color = "#CBD6E2"
            try:
                if first and first.font.color and first.font.color.rgb:
                    color = "#" + str(first.font.color.rgb)
            except Exception:
                pass
            bold = "bold" if (first and first.font.bold) else "normal"
            align = {1: "left", 2: "center", 3: "right"}.get(
                shape.text_frame.paragraphs[0].alignment, "left")
            ls = shape.text_frame.paragraphs[0].line_spacing or 1.15
            html = shape.text_frame.text.replace("&", "&amp;").replace("<", "&lt;")
            parts.append(
                f'<div class="tb {cls}" style="{style}font-size:{size*96/72:.1f}px;'
                f'color:{color};font-weight:{bold};text-align:{align};'
                f'line-height:{ls};font-family:Arial">{html}</div>')
        else:
            fill = "transparent"
            try:
                if shape.fill.type is not None and shape.fill.type == 1:
                    fill = "#" + str(shape.fill.fore_color.rgb)
            except Exception:
                pass
            border = ""
            try:
                if shape.line.color and shape.line.color.rgb:
                    border = f"border:1px solid #{shape.line.color.rgb};"
            except Exception:
                pass
            parts.append(f'<div class="sh {cls}" style="{style}background:{fill};{border}"></div>')
    parts.append("</div>")

OUT.write_text("\n".join(parts))
print(f"wrote {OUT}")
for p in problems:
    print("  !", p)
print(f"{len(problems)} geometry problems")
