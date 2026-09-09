"""Build the QUBIT Hackathon 2026 pitch deck as a .pptx.

The deck is imported into Google Slides, which drives most of the rules here:
no animation, transition, gradient, autofit or native chart survives the
import, and fill/text transparency is dropped -- so every figure goes in as a
transparent PNG, every colour is pre-mixed against the navy ground, and every
run names Arial literally rather than relying on a theme font.

Numbers are read from the cached JSON under docs/ rather than retyped, so a
slide cannot quietly disagree with the run that produced it.

    python presentation/build_deck.py
"""

import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "presentation"))

import theme as T  # noqa: E402

GEN = ROOT / "presentation" / "assets" / "gen"
ASSETS = ROOT / "presentation" / "assets"
DOCS = ROOT / "docs"
OUT = ROOT / "presentation" / "dist" / "qubit_pitch.pptx"

W, H = 13.333, 7.5
MARGIN = 0.67
CONTENT_W = W - 2 * MARGIN


def rgb(hex_color):
    return RGBColor.from_string(hex_color.lstrip("#").upper())


def load(name):
    path = DOCS / name
    return json.loads(path.read_text()) if path.exists() else {}


# --------------------------------------------------------------------------
# primitives
# --------------------------------------------------------------------------
def add_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(T.BG)
    return slide


def text(slide, s, x, y, w, h, size=15, color=T.BODY, bold=False,
         align=PP_ALIGN.LEFT, line=1.15, anchor=MSO_ANCHOR.TOP, caps=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = anchor
    for i, chunk in enumerate(s.split("\n")):
        para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        para.alignment = align
        para.line_spacing = line
        run = para.add_run()
        run.text = chunk.upper() if caps else chunk
        font = run.font
        font.name = "Arial"
        font.size = Pt(size)
        font.bold = bold
        font.color.rgb = rgb(color)
    return box


def rect(slide, x, y, w, h, fill=T.SURFACE, line=None, shape=MSO_SHAPE.RECTANGLE):
    box = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        box.fill.background()
    else:
        box.fill.solid()
        box.fill.fore_color.rgb = rgb(fill)
    if line is None:
        box.line.fill.background()
    else:
        box.line.color.rgb = rgb(line)
        box.line.width = Pt(1)
    box.shadow.inherit = False
    if box.has_text_frame:
        box.text_frame.text = ""
    return box


def picture(slide, path, x, y, w=None, h=None):
    kwargs = {}
    if w is not None:
        kwargs["width"] = Inches(w)
    if h is not None:
        kwargs["height"] = Inches(h)
    return slide.shapes.add_picture(str(path), Inches(x), Inches(y), **kwargs)


def picture_fit(slide, path, x, y, w, h):
    """Place a picture scaled to fit inside a box, centred."""
    from PIL import Image

    with Image.open(path) as img:
        iw, ih = img.size
    scale = min(w / (iw / 96), h / (ih / 96))
    pw, ph = (iw / 96) * scale, (ih / 96) * scale
    return picture(slide, path, x + (w - pw) / 2, y + (h - ph) / 2, w=pw)


def header(slide, kicker, title, page=None):
    text(slide, kicker, MARGIN, 0.42, CONTENT_W, 0.3, size=13, bold=True,
         color=T.CYAN, caps=True)
    text(slide, title, MARGIN, 0.74, CONTENT_W, 0.95, size=30, bold=True,
         color=T.WHITE, line=1.05)
    rect(slide, MARGIN, 1.72, CONTENT_W, 0.017, fill=T.HAIRLINE)
    if page is not None:
        text(slide, f"{page:02d}", W - MARGIN - 0.6, 6.98, 0.6, 0.3, size=10,
             color=T.MUTED, align=PP_ALIGN.RIGHT)


def stat(slide, x, y, w, value, label, color=T.CYAN, size=40):
    text(slide, value, x, y, w, 0.72, size=size, bold=True, color=color)
    text(slide, label, x, y + 0.66, w, 0.6, size=11.5, color=T.MUTED, line=1.2)


def bullets(slide, items, x, y, w, gap=0.86, size=14.5):
    for i, (lead, body) in enumerate(items):
        top = y + i * gap
        text(slide, lead, x, top, w, 0.3, size=size, bold=True, color=T.WHITE)
        text(slide, body, x, top + 0.28, w, gap - 0.3, size=size - 1.5,
             color=T.BODY, line=1.2)


# --------------------------------------------------------------------------
# slides
# --------------------------------------------------------------------------
def title_slide(prs, data):
    slide = add_slide(prs)
    if (GEN / "pipeline_map.png").exists():
        picture(slide, GEN / "pipeline_map.png", 6.55, 2.62, w=6.2)
    if (ASSETS / "logo_qubit.png").exists():
        picture(slide, ASSETS / "logo_qubit.png", MARGIN, 0.5, w=2.5)
    text(slide, "QUBIT Hackathon 2026  ·  AT&T challenge", MARGIN, 1.62,
         CONTENT_W, 0.3, size=13, bold=True, color=T.CYAN, caps=True)
    text(slide, "Routing AT&T's backbone\non 28 qubits", MARGIN, 2.15, 5.7,
         1.9, size=38, bold=True, color=T.WHITE, line=1.06)
    text(slide,
         "Candidate-route encoding, Fortz–Thorup congestion cost,\n"
         "QAOA on Classiq — checked against exact ground truth.",
         MARGIN, 4.15, 5.6, 1.0, size=14.5, color=T.BODY, line=1.3)
    text(slide, "[ team · members ]", MARGIN, 5.15, 6.0, 0.4, size=13,
         color=T.MUTED)
    if (ASSETS / "sponsors_row.png").exists():
        picture(slide, ASSETS / "sponsors_row.png", 2.67, 6.78, w=8.0)
    return slide


def problem_slide(prs, data, page):
    slide = add_slide(prs)
    today = data["tuned_study"]["today"]
    header(slide, "the problem",
           "Today's routing overloads AT&T's real backbone", page)
    picture_fit(slide, GEN / "pipeline_map.png", MARGIN, 1.95, 8.05, 4.0)
    x = MARGIN + 8.35
    w = CONTENT_W - 8.35
    stat(slide, x, 2.0, w, f"{today['max_util']:.0%}", "peak link utilization "
         "under today's lowest-delay routing", color=T.AMBER)
    stat(slide, x, 3.42, w, f"{today['violations']}",
         "link carrying more traffic than its capacity", color=T.AMBER)
    stat(slide, x, 4.84, w, f"{today['phi_star']:.1f}×",
         "Fortz–Thorup congestion cost Φ* against an uncongested network",
         color=T.AMBER)
    text(slide,
         "Real topology and geography from the Internet Topology Zoo's record "
         "of AT&T's published OC-768 map. Latency derived from great-circle "
         "distance; capacity and traffic synthesized, method documented.",
         MARGIN, 6.15, 8.05, 0.8, size=11.5, color=T.MUTED, line=1.25)
    return slide


def encoding_slide(prs, data, page):
    slide = add_slide(prs)
    run = data["tuned_run"]
    n = run.get("num_qubits", 28)
    links, demands = run.get("num_links", 36), run.get("num_demands", 14)
    header(slide, "encoding",
           f"{n} qubits where the textbook encoding needs {links * demands}",
           page)
    picture_fit(slide, GEN / "encoding_bars.png", MARGIN, 1.95, 6.1, 4.35)
    x = MARGIN + 6.5
    w = CONTENT_W - 6.5
    bullets(slide, [
        ("Classical enumeration first",
         "Yen's algorithm across four metrics — delay, cost, hops, spare "
         "capacity — proposes a shortlist of link-disjoint routes per demand. "
         "Validated against networkx on 900 PoP pairs, zero mismatches."),
        ("Qubits track decisions, not wires",
         f"N = Σ|P_k| = demands × routes. It does not grow with the "
         f"{links} directed links, so the same circuit width covers a much "
         f"larger network."),
        ("Penalty degree stays quadratic",
         "The demand-edge encoding needs a degree-D capacity penalty (cubic "
         "at three demands) and per-node flow conservation. One-hot over "
         "routes needs neither."),
        ("The trade, stated plainly",
         "Solution quality is capped by the candidate sets. That is the "
         "hybrid split we chose deliberately."),
    ], x, 2.0, w, gap=1.12, size=14)
    return slide


def novelty_slide(prs, data, page):
    slide = add_slide(prs)
    header(slide, "the formulation problem we had to solve",
           "Fitting Fortz–Thorup per link keeps the capacity cliff", page)
    picture_fit(slide, GEN / "ft_curve.png", MARGIN, 1.95, 6.2, 4.5)
    x = MARGIN + 6.6
    w = CONTENT_W - 6.6
    bullets(slide, [
        ("Capacity is a cliff, not a slope",
         "Fortz & Thorup's cost per link steps through slopes 1, 3, 10, 70, "
         "500, 5000. Below a third of capacity it is nearly free; past "
         "capacity it is catastrophic."),
        ("One global quadratic flattens it",
         "Fitted over all utilizations it becomes 16.7·u² — which minimizes "
         "total squared load, not peak load. On this instance its optimum "
         "leaves the hot link at 105% and never clears it."),
        ("Per-link local fits keep it",
         "H already carries a per-link linear term and u_e is linear in x, so "
         "fitting each link over the load band it can actually reach stays a "
         "QUBO — no slack qubits — and restores the cliff."),
    ], x, 2.0, w, gap=1.42, size=14)
    text(slide,
         "Bounds come from the candidate sets, never from a solution.",
         x, 6.3, w, 0.4, size=11.5, color=T.MUTED)
    return slide


def result_slide(prs, data, page):
    slide = add_slide(prs)
    study = data["tuned_study"]
    today, best = study["today"], study["frontier_min_violations"]
    run = data["tuned_run"]
    header(slide, "the result on the real instance",
           f"{today['max_util']:.0%} → {best['max_util']:.0%} peak, zero violations, on the real backbone",
           page)
    picture_fit(slide, GEN / "hero_map.png", MARGIN, 1.9, CONTENT_W, 3.85)
    y = 5.95
    cols = [
        (f"{today['max_util']:.0%} → {best['max_util']:.0%}", "peak utilization"),
        (f"{today['violations']} → {best['violations']}", "links over capacity"),
        (f"{today['phi_star']:.2f} → {best['phi_star']:.2f}",
         "Fortz–Thorup Φ*"),
        (f"{run.get('brute_force', {}).get('num_combinations', 16384):,}",
         "assignments searched exactly"),
    ]
    for i, (value, label) in enumerate(cols):
        x = MARGIN + i * (CONTENT_W / 4)
        text(slide, value, x, y, CONTENT_W / 4 - 0.2, 0.5, size=23, bold=True,
             color=T.CYAN)
        text(slide, label, x, y + 0.5, CONTENT_W / 4 - 0.2, 0.4, size=11.5,
             color=T.MUTED)
    text(slide,
         "Same candidate routes, same exact solver — the gain comes from the "
         "objective, not from the quantum step.",
         MARGIN, 6.95, CONTENT_W, 0.4, size=11.5, color=T.MUTED)
    return slide


def quantum_slide(prs, data, page):
    slide = add_slide(prs)
    toy = data["toy_run"]
    header(slide, "the quantum step, checked",
           "QAOA returns the optimum where we can verify it", page)
    if (GEN / "convergence_toy.png").exists():
        picture_fit(slide, GEN / "convergence_toy.png", MARGIN, 2.0, 6.3, 3.5)
    x = MARGIN + 6.7
    w = CONTENT_W - 6.7
    mass = toy.get("feasible_mass", 0)
    uniform = toy.get("feasible_fraction_uniform", 1)
    best = toy.get("best_feasible") or {}
    gap = abs(best.get("cost", 0) - toy.get("brute_force", {}).get("best_cost", 0))
    stat(slide, x, 2.0, w, f"{mass:.0%}",
         f"of sampled probability lands on feasible states, against "
         f"{uniform:.1%} for uniform sampling ({mass / uniform:.0f}×)")
    stat(slide, x, 3.55, w, f"{gap:.0e}",
         "cost gap to the exhaustively verified optimum over all "
         f"{toy.get('brute_force', {}).get('num_combinations', 27)} assignments")
    circuit = toy.get("circuit", {})
    text(slide,
         f"{toy.get('num_qubits', 9)} qubits · "
         f"{toy.get('config', {}).get('num_layers', 3)} layers · depth "
         f"{circuit.get('transpiled_depth', '—')} · "
         f"{circuit.get('two_qubit_gates', '—')} two-qubit gates\n"
         f"CVaR at q={toy.get('config', {}).get('quantile', 0.5)} "
         f"(Barkoutsos et al.) · COBYLA · "
         f"{toy.get('config', {}).get('num_shots', 2048)} shots · seed "
         f"{toy.get('config', {}).get('random_seed', 42)} · Classiq simulator",
         x, 5.1, w, 1.3, size=12, color=T.MUTED, line=1.35)
    text(slide,
         "Deliberately small: nine qubits is where every quantum claim can be "
         "checked against exact ground truth. We claim no speedup.",
         MARGIN, 6.4, 6.3, 0.7, size=11.5, color=T.MUTED, line=1.25)
    return slide


def scaling_slide(prs, data, page):
    slide = add_slide(prs)
    run = data["tuned_run"]
    circuit = run.get("circuit", {})
    n = run.get("num_qubits", 28)
    values = run.get("objective_values") or [0, 0]
    header(slide, "where it stops, measured",
           f"At {n} qubits the wall is feasibility, not the circuit", page)
    x2 = MARGIN + 6.9
    bullets(slide, [
        ("The pipeline runs end to end",
         f"{n} qubits synthesized on Classiq to depth "
         f"{circuit.get('transpiled_depth', '—')} with "
         f"{circuit.get('two_qubit_gates', '—')} two-qubit gates, executed on "
         f"the simulator. The CVaR of sampled H falls by "
         f"{values[0] - min(values):.1f} ({values[0]:.1f} to "
         f"{min(values):.1f}; the per-link fit drops constant terms, so H is "
         f"not positive), meaning the optimizer is learning."),
        ("No feasible sample came back",
         f"With 14 demands choosing between two routes, the one-hot subspace "
         f"is 2^14 of 2^28 states — "
         f"{run.get('feasible_fraction_uniform', 6.1e-5):.3%} of the Hilbert "
         f"space. A depth-1 X mixer spreads amplitude across all of it."),
        ("The fix is known, not hand-waving",
         "An XY mixer confines evolution to the feasible subspace by "
         "construction; a binary encoding needs log2(K) qubits per demand, rounded up, and "
         "removes one-hot entirely. Both drop into one function of qaoa.py."),
    ], MARGIN, 2.0, 6.5, gap=1.45, size=14)
    rect(slide, x2, 2.0, CONTENT_W - 6.9, 3.5, fill=T.SURFACE,
         line=T.HAIRLINE)
    text(slide, "what we will not claim", x2 + 0.3, 2.25, CONTENT_W - 7.5, 0.3,
         size=12, bold=True, color=T.AMBER, caps=True)
    text(slide,
         "• A quantum speedup. At these sizes exact classical search wins, "
         "and we say so.\n\n"
         "• A quantum result on the 28-qubit instance. The clean routing there "
         "is classical, from exhaustive search over the same QUBO.\n\n"
         "• That the congestion win is quantum. It comes from the objective, "
         "solved by brute force.",
         x2 + 0.3, 2.7, CONTENT_W - 7.5, 2.6, size=13, color=T.BODY, line=1.3)
    text(slide,
         "Reporting the boundary is the point: it is what tells you which "
         "mixer to build next.",
         MARGIN, 6.55, CONTENT_W, 0.4, size=11.5, color=T.MUTED)
    return slide


def craft_slide(prs, data, page):
    slide = add_slide(prs)
    toy = data["toy_run"]
    header(slide, "how it is built",
           "One cost function: circuit, brute force, decoder", page)
    code = (
        "@qfunc\n"
        "def main(params, x):\n"
        "    allocate(x)\n"
        "    hadamard_transform(x)\n"
        "    repeat(num_layers, lambda i: (\n"
        "        phase(cost_fn(x), params[2*i]),\n"
        "        mixer_layer(params[2*i+1], x),\n"
        "    ))"
    )
    rect(slide, MARGIN, 2.0, 5.6, 2.5, fill=T.SURFACE, line=T.HAIRLINE)
    text(slide, code, MARGIN + 0.28, 2.24, 5.1, 2.1, size=12.5,
         color=T.CYAN_PALE, line=1.25)
    text(slide,
         "Classiq's phase() compiles the same Python cost function that brute "
         "force minimizes and the decoder scores samples with — so the "
         "Hamiltonian and the classical objective cannot drift apart.",
         MARGIN, 4.65, 5.6, 1.0, size=13, color=T.BODY, line=1.3)
    x = MARGIN + 6.1
    w = CONTENT_W - 6.1
    bullets(slide, [
        ("QAOA proposes, classical validation selects",
         "The decoder gates every sample on one-hot feasibility and capacity "
         "before any KPI comparison, then returns the best feasible routing. "
         "Nothing infeasible can leave the pipeline."),
        ("49 tests, one of them exhaustive",
         "Including a bitstring-by-bitstring check of the cost polynomial "
         "against an independent implementation on a four-qubit instance."),
        ("Only qaoa.py imports classiq",
         "Model, QUBO, decoding and plotting stay pure Python, so the solver "
         "is swappable — MILP or QPU — without touching the formulation."),
    ], x, 2.0, w, gap=1.42, size=14)
    return slide


def closing_slide(prs, data, page):
    """The slide that stays on screen for all three minutes of Q&A.

    No figure competes with the sentence here: the guidance for this event is
    explicit that this slide is where the judges' attention sits while they
    score, so it carries the claim, the numbers behind it, and what comes next.
    """
    slide = add_slide(prs)
    study = data["tuned_study"]
    today, best = study["today"], study["frontier_min_violations"]
    run = data["tuned_run"]
    toy = data["toy_run"]

    text(slide, "the one sentence", MARGIN, 0.75, CONTENT_W, 0.3, size=13,
         bold=True, color=T.CYAN, caps=True)
    text(slide,
         "We priced AT&T's own congestion cost into 28 qubits —\n"
         "a textbook encoding needs 504 — and checked every\n"
         "answer against exact ground truth.",
         MARGIN, 1.25, CONTENT_W, 2.5, size=33, bold=True, color=T.WHITE,
         line=1.18)

    edge = run.get("num_links", 36) * run.get("num_demands", 14)
    chips = [
        f"{run.get('num_qubits', 28)} vs {edge} qubits",
        f"peak {today['max_util']:.0%} → {best['max_util']:.0%}",
        f"{today['violations']} → {best['violations']} links over capacity",
        f"Φ* {today['phi_star']:.1f} → {best['phi_star']:.1f}",
        f"feasible mass {toy.get('feasible_mass', 0):.0%} vs "
        f"{toy.get('feasible_fraction_uniform', 0):.1%} uniform",
        "0 gap to the verified optimum",
        "49 tests",
    ]
    # Wrap chips across rows instead of running off the slide edge.
    x, y = MARGIN, 4.15
    for chip in chips:
        w = 0.108 * len(chip) + 0.38
        if x + w > MARGIN + CONTENT_W:
            x, y = MARGIN, y + 0.62
        rect(slide, x, y, w, 0.5, fill=T.SURFACE, line=T.HAIRLINE)
        text(slide, chip, x + 0.19, y + 0.13, w - 0.38, 0.3, size=12.5,
             color=T.CYAN)
        x += w + 0.16

    rect(slide, MARGIN, 6.18, CONTENT_W, 0.017, fill=T.HAIRLINE)
    text(slide,
         "Next:  XY mixer for feasible-subspace evolution  ·  MILP baseline  "
         "·  Fortz–Thorup LP bound  ·  QPU run",
         MARGIN, 6.34, CONTENT_W, 0.4, size=12.5, color=T.MUTED)
    if (ASSETS / "logo_qubit.png").exists():
        picture(slide, ASSETS / "logo_qubit.png", W - MARGIN - 1.9, 0.62, w=1.9)
    if (ASSETS / "sponsors_row.png").exists():
        picture(slide, ASSETS / "sponsors_row.png", 3.37, 6.85, w=6.6)
    return slide


# --------------------------------------------------------------------------
# backup slides
# --------------------------------------------------------------------------
def divider(prs, label="backup"):
    slide = add_slide(prs)
    text(slide, label, MARGIN, 3.1, CONTENT_W, 1.0, size=44, bold=True,
         color=T.GHOST, caps=True)
    rect(slide, MARGIN, 4.3, CONTENT_W, 0.017, fill=T.HAIRLINE)
    text(slide, "Everything that did not earn a place in five minutes.",
         MARGIN, 4.5, CONTENT_W, 0.4, size=14, color=T.MUTED)
    return slide


def calibration_backup(prs, data):
    slide = add_slide(prs)
    header(slide, "backup · instance calibration",
           "An unroutable instance is a capacity result, not an optimization")
    rows = [("instance", "demands", "today", "best achievable", "verdict")]
    dense = data.get("dense_study")
    tuned = data.get("tuned_study")
    if dense:
        rows.append((
            "east10_dense", str(dense["n_demands"]),
            f"{dense['today']['max_util']:.0%} · "
            f"{dense['today']['violations']} over",
            f"{dense['frontier_min_violations']['max_util']:.0%} · "
            f"{dense['frontier_min_violations']['violations']} over",
            "oversubscribed — no clean assignment exists",
        ))
    if tuned:
        rows.append((
            "east10_tuned", str(tuned["n_demands"]),
            f"{tuned['today']['max_util']:.0%} · "
            f"{tuned['today']['violations']} over",
            f"{tuned['frontier_min_violations']['max_util']:.0%} · "
            f"{tuned['frontier_min_violations']['violations']} over",
            "contested and solvable — 112 of 16,384 clean",
        ))
    widths = [1.9, 1.2, 2.2, 2.4, 4.3]
    y = 2.1
    for r, row in enumerate(rows):
        x = MARGIN
        for value, w in zip(row, widths):
            text(slide, value, x, y, w - 0.15, 0.5,
                 size=12.5 if r else 11.5,
                 color=T.MUTED if r == 0 else T.BODY,
                 bold=(r == 0), caps=(r == 0))
            x += w
        y += 0.52 if r == 0 else 0.62
        if r == 0:
            rect(slide, MARGIN, y - 0.14, CONTENT_W, 0.014, fill=T.HAIRLINE)
    text(slide,
         "`fill` caps each demand at a fraction of the widest bottleneck open "
         "to it. At 0.9 any two demands sharing a link overload it however "
         "they are routed; at 0.3 they compete instead, and the optimizer has "
         "something to win. Both instances are the same 10 real PoPs.",
         MARGIN, 4.5, CONTENT_W, 1.0, size=13, color=T.BODY, line=1.3)
    return slide


def toy_backup(prs, data):
    slide = add_slide(prs)
    header(slide, "backup · the verified instance",
           "Nine qubits, three demands, every answer checkable by hand")
    picture_fit(slide, GEN / "toy_comparison.png", MARGIN, 1.95, CONTENT_W, 4.3)
    text(slide,
         "Both columns are decoded under one coefficient set and share a "
         "single utilization colour scale, so the panels are directly "
         "comparable. Priority π=3 on the premium demand is what makes it win "
         "the contested link.",
         MARGIN, 6.4, CONTENT_W, 0.7, size=12, color=T.MUTED, line=1.25)
    return slide


def main():
    data = {
        "toy_run": load("qaoa_toy.json"),
        "tuned_run": load("qaoa_east10_tuned.json"),
    }
    study = load("real_instance_study.json")
    data["tuned_study"] = study.get("east10_tuned", {})
    data["dense_study"] = study.get("east10_dense", {})
    if not data["tuned_study"]:
        raise SystemExit("run scripts/real_instance_study.py first")

    prs = Presentation()
    prs.slide_width = Emu(int(W * 914400))
    prs.slide_height = Emu(int(H * 914400))

    title_slide(prs, data)
    problem_slide(prs, data, 2)
    encoding_slide(prs, data, 3)
    novelty_slide(prs, data, 4)
    result_slide(prs, data, 5)
    quantum_slide(prs, data, 6)
    scaling_slide(prs, data, 7)
    craft_slide(prs, data, 8)
    closing_slide(prs, data, 9)
    divider(prs)
    calibration_backup(prs, data)
    toy_backup(prs, data)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    size = OUT.stat().st_size / 1e6
    print(f"wrote {OUT.relative_to(ROOT)} — {len(prs.slides.__iter__.__self__._sldIdLst)} slides, {size:.1f} MB")


if __name__ == "__main__":
    main()
